"""Tests for DOCX style-pack integration."""

import docx
import pytest
from docx.opc.exceptions import PackageNotFoundError
from pathlib import Path
import shutil

from dkit.doc2 import builder
from dkit.doc2 import document as doc
from dkit.doc2.document import Document
from dkit.doc2.docx_helper import DocxConfig
from dkit.doc2.docx_renderer import DocxRenderer
from dkit.stylepack.matplotlib import matplotlib_theme

from test_stylepack_reportlab import load_pack


def make_template(pack):
    """Create a valid template with visible named-style content."""
    template = docx.Document()
    template.styles["Normal"].font.name = "Arial"
    template.styles["Heading 1"].font.name = "Arial"
    template.sections[0].header.paragraphs[0].text = "Template header"
    template.sections[0].footer.paragraphs[0].text = "Template footer"
    template.add_paragraph("Template title page", style="Title")
    template.save(pack.root / "docx" / "template.docx")


def test_style_pack_passes_template_and_preserves_named_styles(tmp_path):
    """The factory opens the registered template and uses its style system."""
    pack = load_pack(tmp_path)
    make_template(pack)
    document = Document(title="Example")
    document.add_template("# Heading\n\nBody text")

    renderer = builder.get_renderer(document, "docx", style_pack=pack)
    output = tmp_path / "output.docx"
    renderer.render(str(output))
    result = docx.Document(output)

    assert result.sections[0].header.paragraphs[0].text == "Template header"
    assert result.sections[0].footer.paragraphs[0].text == "Template footer"
    assert any(p.style.name == "Title" for p in result.paragraphs)
    assert any(p.style.name == "Heading 1" for p in result.paragraphs)
    assert any(p.style.name == "Normal" for p in result.paragraphs)


def test_style_pack_selects_its_docx_table_style(tmp_path):
    """A DOCX manifest can select a named table style from its template."""
    pack = load_pack(tmp_path, table_style="Future Blue table")
    reference = Path(__file__).parents[1] / "reference_template.docx"
    shutil.copyfile(reference, pack.root / "docx" / "template.docx")
    document = Document(title="Example")
    document.add_element(
        doc.Table(
            data=[{"name": "Example", "value": 1}],
            columns=[
                doc.Column("name", "Name", width=4),
                doc.Column("value", "Value", width=4),
            ],
        )
    )

    renderer = builder.get_renderer(
        document, "docx", style_pack=pack
    )
    output = tmp_path / "output.docx"
    renderer.render(str(output))
    result = docx.Document(output)

    assert result.tables[-1].style.name == "Future Blue table"


def test_docx_renderer_keeps_configured_table_style_without_style_pack():
    """The existing DocxConfig table style remains the fallback."""
    renderer = DocxRenderer(
        Document(title="Example"),
        config=DocxConfig(sty_table="Table Grid"),
    )

    assert renderer.table_style == "Table Grid"


@pytest.mark.parametrize(
    ("size", "orientation", "expected"),
    [
        ("a4", "portrait", (21.0, 29.7)),
        ("letter", "landscape", (27.94, 21.59)),
        ("legal", "portrait", (21.59, 35.56)),
        ("a5", "landscape", (21.0, 14.8)),
    ],
)
def test_style_pack_sets_section_page_geometry(
    tmp_path, size, orientation, expected
):
    """DOCX section dimensions follow the declared physical page tokens."""
    pack = load_pack(tmp_path / f"{size}-{orientation}", size, orientation)
    make_template(pack)
    renderer = builder.get_renderer(
        Document(title="Example"), "docx", style_pack=pack
    )

    section = renderer.xdoc.sections[0]
    assert round(section.page_width.cm, 2) == expected[0]
    assert round(section.page_height.cm, 2) == expected[1]
    assert round(section.left_margin.cm, 2) == 2.0
    assert round(section.top_margin.cm, 2) == 2.2


def test_invalid_style_pack_template_fails_before_output(tmp_path):
    """A malformed registered template cannot create a partial DOCX."""
    pack = load_pack(tmp_path)
    output = tmp_path / "output.docx"

    with pytest.raises(PackageNotFoundError):
        builder.render_to_file(
            Document(title="Example"),
            "docx",
            str(output),
            style_pack=pack,
        )

    assert not output.exists()


def test_docx_build_uses_print_plot_variant(tmp_path, monkeypatch):
    """Styled DOCX builds select the print matplotlib asset."""
    pack = load_pack(tmp_path)
    make_template(pack)
    definition = builder.DocumentDefinition(
        info=builder.DocumentInfo(title="Example", contact=""),
        configuration=builder.DocumentConfiguration(
            renderer="docx", style="dkit-blue", output="unused.docx"
        ),
        templates=[],
        code={},
        data={},
        variables={},
    )
    observed = []
    monkeypatch.setattr(builder, "resolve_style", lambda _name: pack)
    monkeypatch.setattr(
        builder,
        "matplotlib_theme",
        lambda _pack, variant: observed.append(variant) or matplotlib_theme(
            _pack, variant
        ),
    )
    monkeypatch.setattr(builder, "render_to_file", lambda *args, **kwargs: None)

    builder.Builder(definition).build()

    assert observed == ["print"]
